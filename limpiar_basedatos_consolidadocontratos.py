import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import tempfile
import os
import warnings
import re

warnings.filterwarnings('ignore')

# ─────────────────────────────────────────────────────────────
# ✅ LISTA MAESTRA ESTRICTA DE COMPRADORES
# ─────────────────────────────────────────────────────────────
STRATEGIC_BUYERS = {
    'Patricio Espinoza', 'Jorge Urrutia', 'Bárbara García', 'Claudio Berrios',
    'Martina Fuentes', 'Joseph España', 'Michelle Palma', 'Juan Figueroa',
    'Magdalena Farias', 'Denisse Andrea Gonzalez Terrile'
}

TACTICAL_BUYERS = {
    'Leonardo Nacarate', 'Martina Fuentes', 'Scarlette Lucero',
    'Margarita Lineros', 'Erika Silva', 'Karina Satelo', 'Pablo Labs'
}

TYPO_CORRECTIONS = {
    'jorge uturria': 'Jorge Urrutia', 'jorgue urrutia': 'Jorge Urrutia',
    'dennis andrea gonzales': 'Denisse Andrea Gonzalez Terrile',
    'denisse andrea gonzalez terrile': 'Denisse Andrea Gonzalez Terrile',
    'juan daniel figueroa': 'Juan Figueroa',
    'joseph eduardo españa escalona': 'Joseph España',
    'michelle esperanza': 'Michelle Palma',
    'leonardo nacarete': 'Leonardo Nacarate',
    'martina fuentes': 'Martina Fuentes'
}

def normalize_name(name: str) -> str:
    if pd.isna(name) or str(name).strip() == '': return ''
    clean = str(name).strip().lower()
    clean = ''.join(c for c in clean if c not in 'áéíóúüñ')
    return clean

def classify_buyer_strict(raw_name: str) -> tuple:
    clean_raw = normalize_name(raw_name)
    if not clean_raw: return None, None
    for typo, correct in TYPO_CORRECTIONS.items():
        if typo in clean_raw or clean_raw in typo:
            clean_raw = normalize_name(correct); break
    for official in STRATEGIC_BUYERS:
        if clean_raw == normalize_name(official) or clean_raw in normalize_name(official) or normalize_name(official) in clean_raw:
            return 'strategic', official
    for official in TACTICAL_BUYERS:
        if clean_raw == normalize_name(official) or clean_raw in normalize_name(official) or normalize_name(official) in clean_raw:
            return 'tactical', official
    return None, None

# ─────────────────────────────────────────────────────────────
# 📐 ESTRUCTURA DEL CONSOLIDADO (30 Columnas)
# ─────────────────────────────────────────────────────────────
TARGET_HEADERS = [
    'Contrato Sap', 'Contrato Legado', 'Comprador Estratégico', 'Comprador Táctico',
    'Estado Contrato Ariba', 'Rut', 'Cód SAP', 'Proveedor', 'Fecha Inicio',
    'Fecha Término Contrato', 'Estado Contrato', 'Descripción', 'Área', 'Gerencia',
    'Planta', 'Ingresa a Planta', 'Aplica Boleta de Garantía (Ariba)',
    'Aplica Boleta de Garantía (Contrato firmado)', 'Tipo Garantía', 'N° Garantia',
    'Moneda Garantía', 'Monto Garantía', 'Vencimiento Garantía', 'Estado Garantía',
    'Administrador de Contrato', 'Correo Electrónico', 'Observación contrato Control Contratista',
    'Observación Control Contratistas Boleta de Garantía', 'Contratos Indefinidos', 'Observación Interna'
]

# 🔄 MAPEO: Campos técnicos de Ariba → Columnas del Consolidado
ARIBA_TO_CONSOLIDADO = {
    'ContractId': 'Contrato Sap',
    'ProjectInfo.ProjectName': 'Contrato Legado',
    'Owner.UserName': 'Comprador Estratégico',
    'ContractStatus': 'Estado Contrato Ariba',
    'UF_string11': 'Rut',
    'UF_string10': 'Cód SAP',
    'AffectedParties.CommonSupplierName': 'Proveedor',
    'EffectiveDate.Day': 'Fecha Inicio',
    'ExpirationDate.Day': 'Fecha Término Contrato',
    'Description': 'Descripción',
    'Region.RegionNameL2': 'Área',
    'IsEvergreen': 'Contratos Indefinidos',
    'UF_boolean1': 'Aplica Boleta de Garantía (Ariba)',
    'UF_string23': 'Tipo Garantía',
    'UF_time6.Day': 'Vencimiento Garantía',
    'sum(Amount)': 'Monto Garantía',
}

STYLES = {
    "header": {
        "fill": PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid"),
        "font": Font(name='Arial', size=8, bold=True),
        "alignment": Alignment(horizontal='center', vertical='center', wrap_text=True),
        "border": Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin'))
    },
    "data": {
        "font": Font(name='Arial', size=8),
        "border": Border(left=Side('thin'), right=Side('thin'), top=Side('thin'), bottom=Side('thin')),
        "alignment": Alignment(horizontal='left', vertical='center', wrap_text=False)
    }
}

COLUMN_WIDTHS = [14, 16, 20, 16, 20, 16, 12, 38, 14, 18, 16, 45, 16, 16, 12, 14, 26, 32, 14, 14, 14, 16, 16, 14, 20, 22, 35, 45, 18, 45]

def load_pivot_ariba(file_path: str) -> pd.DataFrame:
    """
    Carga el Pivot de Ariba con formato especial:
    - Fila con 'Raw_Field_Names' tiene todos los headers concatenados por coma
    - Los datos reales vienen en filas posteriores
    """
    # Leer primeras filas para detectar estructura
    df_raw = pd.read_excel(file_path, header=None, nrows=100)
    
    # Buscar fila con 'Raw_Field_Names'
    header_row_idx = None
    for i, row in df_raw.iterrows():
        row_str = ' '.join(str(v).lower() for v in row if pd.notna(v))
        if 'raw_field_names' in row_str:
            header_row_idx = i
            break
    
    if header_row_idx is None:
        raise ValueError("No se encontró la fila 'Raw_Field_Names' con los encabezados del Pivot.")
    
    # Extraer y parsear los nombres de columnas desde la celda concatenada
    raw_headers_cell = df_raw.iloc[header_row_idx, 2]  # Columna C típicamente
    if pd.isna(raw_headers_cell):
        # Intentar en otras columnas
        for col in range(df_raw.shape[1]):
            val = df_raw.iloc[header_row_idx, col]
            if pd.notna(val) and 'ContractId' in str(val):
                raw_headers_cell = val
                break
    
    # Separar por coma y limpiar
    column_names = [c.strip() for c in str(raw_headers_cell).split(',') if c.strip()]
    
    # Leer los datos reales (saltando filas de metadata)
    data_start_row = header_row_idx + 1
    df_data = pd.read_excel(file_path, header=None, skiprows=range(data_start_row), engine='openpyxl')
    
    # Asignar nombres de columna
    if len(column_names) <= df_data.shape[1]:
        df_data.columns = column_names + [f'Unnamed_{i}' for i in range(len(column_names), df_data.shape[1])]
    else:
        df_data.columns = column_names[:df_data.shape[1]]
    
    # Limpiar filas vacías
    df_data = df_data.dropna(how='all').reset_index(drop=True)
    
    return df_data

def transform_data(df_pivot: pd.DataFrame) -> tuple:
    """Transforma datos del Pivot Ariba al formato Consolidado."""
    df_out = pd.DataFrame(columns=TARGET_HEADERS)
    
    # 1. Mapeo de campos Ariba → Consolidado
    for ariba_col, target_col in ARIBA_TO_CONSOLIDADO.items():
        if ariba_col in df_pivot.columns and target_col in df_out.columns:
            df_out[target_col] = df_pivot[ariba_col].copy()
    
    # Duplicar estado si existe
    if 'Estado Contrato Ariba' in df_out.columns and 'Estado Contrato' in df_out.columns:
        df_out['Estado Contrato'] = df_out['Estado Contrato Ariba']
    
    # 2. 🔥 VALIDACIÓN ESTRICTA DE COMPRADORES
    if 'Owner.UserName' in df_pivot.columns:
        raw_owners = df_pivot['Owner.UserName'].fillna('').astype(str)
        classified = raw_owners.apply(classify_buyer_strict)
        df_out['Comprador Estratégico'] = [x[1] if x[0] == 'strategic' else '' for x in classified]
        df_out['Comprador Táctico'] = [x[1] if x[0] == 'tactical' else '' for x in classified]
    
    # 🗑️ Eliminar filas sin compradores válidos
    mask_valid_buyer = (df_out['Comprador Estratégico'] != '') | (df_out['Comprador Táctico'] != '')
    dropped_invalid = (~mask_valid_buyer).sum()
    df_out = df_out[mask_valid_buyer].reset_index(drop=True)
    
    # 🚫 Eliminar contratos Cerrados
    if 'Estado Contrato Ariba' in df_out.columns:
        mask_no_cerrado = ~df_out['Estado Contrato Ariba'].astype(str).str.strip().str.lower().isin(['cerrado', 'cerrados'])
        dropped_cerrados = (~mask_no_cerrado).sum()
        df_out = df_out[mask_no_cerrado].reset_index(drop=True)
    else:
        dropped_cerrados = 0
    
    # 🛡️ Copiar Estratégico a Táctico si está vacío
    mask_tactical_empty = (df_out['Comprador Táctico'] == '') & (df_out['Comprador Estratégico'] != '')
    df_out.loc[mask_tactical_empty, 'Comprador Táctico'] = df_out.loc[mask_tactical_empty, 'Comprador Estratégico']
    
    # Administrador = Estratégico por defecto
    df_out['Administrador de Contrato'] = df_out['Comprador Estratégico'].where(
        df_out['Comprador Estratégico'] != '', df_out['Comprador Táctico']
    )
    
    # 3. Formateo de fechas
    for date_col in ['Fecha Inicio', 'Fecha Término Contrato', 'Vencimiento Garantía']:
        if date_col in df_out.columns:
            df_out[date_col] = pd.to_datetime(df_out[date_col], errors='coerce').dt.strftime('%d-%m-%Y')
            df_out.loc[df_out[date_col] == 'NaT', date_col] = ''
    
    # 4. Normalizar Sí/No
    for col in ['Ingresa a Planta', 'Aplica Boleta de Garantía (Ariba)', 'Aplica Boleta de Garantía (Contrato firmado)', 'Contratos Indefinidos']:
        if col in df_out.columns:
            df_out[col] = df_out[col].astype(str).str.strip().str.title()
            df_out[col] = df_out[col].replace(['Si', 'Sí', 'Yes', 'Y', 'True'], 'Sí')
            df_out[col] = df_out[col].replace(['No', 'N', 'False', 'Nan', ''], 'No')
    
    # 5. Limpieza final
    df_out = df_out.fillna('')
    df_out = df_out.replace(['null', 'None', 'Unclassified', 'nan'], '')
    
    return df_out, dropped_invalid, dropped_cerrados

def apply_formatting(df: pd.DataFrame, output_path: str):
    """Aplica formato Excel idéntico al Consolidado de referencia."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado de Contratos"
    
    for col_idx, header in enumerate(TARGET_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.fill = STYLES["header"]["fill"]
        cell.font = STYLES["header"]["font"]
        cell.alignment = STYLES["header"]["alignment"]
        cell.border = STYLES["header"]["border"]
    ws.row_dimensions[1].height = 45
    
    for r_idx, row in df.iterrows():
        for c_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=r_idx + 2, column=c_idx, value=val)
            cell.font = STYLES["data"]["font"]
            cell.border = STYLES["data"]["border"]
            cell.alignment = STYLES["data"]["alignment"]
            if c_idx in [9, 10, 23]:
                cell.alignment = Alignment(horizontal='center', vertical='center')
            elif isinstance(val, (int, float)):
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.number_format = '#,##0.00'
                
    for i, w in enumerate(COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    
    ws.auto_filter.ref = f"A1:{get_column_letter(len(TARGET_HEADERS))}{len(df) + 1}"
    ws.freeze_panes = "C2"
    wb.save(output_path)

# ─────────────────────────────────────────────────────────────
# 🌐 STREAMLIT UI
# ─────────────────────────────────────────────────────────────
st.set_page_config(page_title="Generador Consolidado Ariba", layout="centered")
st.title("📑 Generador de Consolidado de Contratos")
st.caption("Sube el Pivot de Ariba. Se parsearán los campos técnicos correctamente.")

uploaded_file = st.file_uploader("📥 Archivo Pivot (.xlsx)", type=["xlsx"])

if uploaded_file:
    with st.spinner("Parseando Pivot de Ariba, validando compradores y aplicando formato..."):
        try:
            with tempfile.NamedTemporaryFile(delete=False, suffix="_pivot.xlsx") as tmp:
                tmp.write(uploaded_file.getvalue())
                pivot_path = tmp.name
            
            df_pivot = load_pivot_ariba(pivot_path)
            st.info(f"📊 Pivot cargado: {len(df_pivot)} filas, columnas detectadas: {len([c for c in df_pivot.columns if 'Unnamed' not in c])}")
            
            df_final, dropped_invalid, dropped_cerrados = transform_data(df_pivot)
            
            out_path = pivot_path.replace("_pivot.xlsx", "_Consolidado_Final.xlsx")
            apply_formatting(df_final, out_path)
            
            with open(out_path, "rb") as f:
                st.download_button(
                    label="📥 Descargar Consolidado de Contratos",
                    data=f,
                    file_name="Consolidado de Contratos.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True
                )
            
            st.success(f"✅ **Archivo generado.**\n• 📄 {len(df_final)} contratos procesados.\n• 🗑️ {dropped_invalid} eliminados por compradores no oficiales.\n• 🚫 {dropped_cerrados} eliminados por estado 'Cerrado'.")
            
            os.unlink(pivot_path)
            os.unlink(out_path)
            
        except Exception as e:
            st.error(f"❌ Error: {str(e)}")
            st.exception(e)
